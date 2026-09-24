"""Fill the company Flavor Sheet template from a flavor sheet's data.

Fill, never regenerate: copy the template, write inputs into known cells
(docs/phase_f_templates_layout.md §2), keep every style. The g column gets the
template's own formula so Excel does the arithmetic. A flavor longer than its
slot grows the block with inserted rows; every four flavors past the first
four go on a copy of the template sheet. Needs openpyxl; the ETL never imports it.
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .db import PROJECT_ROOT

DEFAULT_TEMPLATE_PATH = PROJECT_ROOT / "data" / "real" / "templates" / "Flavor Sheet Blank Template.xlsx"

FLAVORS_PER_PAGE = 4
LAST_COLUMN = 8  # A..H — the printed width of the form

# Template geometry (Sheet1). Each block is two slots side by side sharing
# rows: first_row is the BASE row, last_row the thick-bordered bottom row.
TOP_TITLE_ROW, TOP_FIRST_ROW, TOP_LAST_ROW = 6, 8, 20
BOTTOM_TITLE_ROW, BOTTOM_FIRST_ROW, BOTTOM_LAST_ROW = 21, 23, 34
# (name column, mg column, g column, title column) for the left and right slot.
SLOT_COLUMNS = (("A", "B", "C", "B"), ("E", "F", "G", "F"))


@dataclass
class Flavor:
    """One flavor slot as printed: the title, the BASE mg, and (name, mg) lines."""
    title: str
    base_mg: float | None = None
    lines: list[tuple[str, float | None]] = field(default_factory=list)


def render(
    product_line: str,
    servings: float | None,
    flavors: list[Flavor],
    template_path: Path | str = DEFAULT_TEMPLATE_PATH,
) -> bytes:
    """Return the filled workbook as .xlsx bytes, four flavors per page."""
    wb = load_workbook(template_path)
    template = wb.worksheets[0]

    pages = [flavors[i:i + FLAVORS_PER_PAGE] for i in range(0, len(flavors), FLAVORS_PER_PAGE)] or [[]]
    sheets = [template]
    for n in range(1, len(pages)):
        # Copy before the first page is filled, so every page starts blank.
        extra = wb.copy_worksheet(template)
        first = n * FLAVORS_PER_PAGE + 1
        extra.title = f"Flavors {first}-{first + FLAVORS_PER_PAGE - 1}"
        wb.move_sheet(extra, offset=wb.index(sheets[-1]) + 1 - wb.index(extra))
        sheets.append(extra)

    for ws, page in zip(sheets, pages):
        _fill_page(ws, product_line, servings, page)

    wb.calculation.fullCalcOnLoad = True
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _fill_page(ws: Worksheet, product_line: str, servings: float | None, page: list[Flavor]) -> None:
    """Fill one page: the header cells, then the top and bottom pairs of flavor slots."""
    # H3/F4/H4 aren't part of the sheet R&D fills today; clear the template's placeholders.
    ws["A3"] = f"Product Name: {product_line}"
    ws["F3"] = servings
    ws["H3"] = None
    ws["F4"] = None
    ws["H4"] = None

    top, bottom = page[:2], page[2:4]

    top_extra = _rows_needed(top) - (TOP_LAST_ROW - TOP_FIRST_ROW + 1)
    if top_extra > 0:
        _grow(ws, TOP_LAST_ROW, top_extra)
    top_extra = max(top_extra, 0)
    _fill_block(ws, TOP_TITLE_ROW, TOP_FIRST_ROW, TOP_LAST_ROW + top_extra, top)

    b_title, b_first, b_last = (r + top_extra for r in (BOTTOM_TITLE_ROW, BOTTOM_FIRST_ROW, BOTTOM_LAST_ROW))
    bottom_extra = _rows_needed(bottom) - (b_last - b_first + 1)
    if bottom_extra > 0:
        _grow(ws, b_last, bottom_extra)
    _fill_block(ws, b_title, b_first, b_last + max(bottom_extra, 0), bottom)


def _rows_needed(pair: list[Flavor]) -> int:
    """Return the rows a block must hold: BASE plus the longer flavor's lines."""
    return 1 + max((len(f.lines) for f in pair), default=0)


def _fill_block(ws: Worksheet, title_row: int, first_row: int, last_row: int, pair: list[Flavor]) -> None:
    """Clear a block's example data and write up to two flavors side by side."""
    for slot, (name_col, mg_col, g_col, title_col) in enumerate(SLOT_COLUMNS):
        # Clear the template's example data. The BASE label (first row, name
        # column) is static template text and stays.
        for r in range(first_row, last_row + 1):
            if r != first_row:
                ws[f"{name_col}{r}"] = None
            ws[f"{mg_col}{r}"] = None
            ws[f"{g_col}{r}"] = None

        flavor = pair[slot] if slot < len(pair) else None
        ws[f"{title_col}{title_row}"] = flavor.title if flavor else None
        if flavor is None:
            continue

        _write_amount(ws, mg_col, g_col, first_row, flavor.base_mg)
        for offset, (name, mg) in enumerate(flavor.lines, start=1):
            ws[f"{name_col}{first_row + offset}"] = name
            _write_amount(ws, mg_col, g_col, first_row + offset, mg)


def _write_amount(ws: Worksheet, mg_col: str, g_col: str, row: int, mg: float | None) -> None:
    """Write mg and the g formula for one row; a blank mg leaves both blank rather than print 0.000 g."""
    if mg is None:
        return
    ws[f"{mg_col}{row}"] = mg
    g = ws[f"{g_col}{row}"]
    g.value = f"={mg_col}{row}*$F$3/1000"
    g.number_format = "0.000"


def _grow(ws: Worksheet, before: int, count: int) -> None:
    """Insert `count` styled rows above row `before` (a block's last row), keeping the banding.

    openpyxl's insert_rows moves neither merged ranges nor row heights, so both
    are shifted here.
    """
    heights = {
        r: ws.row_dimensions[r].height
        for r in range(before, ws.max_row + 1)
        if ws.row_dimensions[r].height is not None
    }
    ws.insert_rows(before, count)

    for merged in ws.merged_cells.ranges:
        if merged.min_row >= before:
            merged.shift(0, count)

    for r in list(heights):
        ws.row_dimensions[r].height = None
    for r, h in heights.items():
        ws.row_dimensions[r + count].height = h

    for r in range(before, before + count):
        for c in range(1, LAST_COLUMN + 1):
            ws.cell(r, c)._style = copy(ws.cell(r - 2, c)._style)

    last = before + count
    for c in range(1, LAST_COLUMN + 1):
        ws.cell(last, c).fill = copy(ws.cell(last - 2, c).fill)
