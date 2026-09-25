"""Fill the company Sample Record Sheet template from a flavor's data.

Fill, never regenerate: write only the header and per-line input cells
(docs/phase_f_templates_layout.md §1); every computed column and total is the
sheet's own formula. The sections have fixed capacity, so an oversized formula
is refused (RecordSheetOverflow) rather than truncated, and a missing price is
left blank, never 0. Needs openpyxl; the ETL never imports it.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from .db import PROJECT_ROOT

DEFAULT_TEMPLATE_PATH = (
    PROJECT_ROOT / "data" / "real" / "templates" / "Sample Record Sheet Blank Template.xlsx"
)

# Template geometry (Sheet1), from docs/phase_f_templates_layout.md §1.
ACTIVE_FIRST_ROW, ACTIVE_CAPACITY = 12, 23        # rows 12-34, subtotal row 35
EXCIPIENT_FIRST_ROW, EXCIPIENT_CAPACITY = 37, 12  # rows 37-48, subtotal row 49

# The input columns this renderer owns. F/G/H/J/K are the sheet's formulas and
# are left untouched; clearing an example only touches these.
_LINE_INPUT_COLUMNS = ("A", "B", "C", "D", "E", "I")

# Header input cell -> attribute on RecordSheet.
_HEADER_CELLS = {
    "B3": "brand",
    "B4": "product",
    "B5": "flavor",
    "B6": "sample_code",
    "B7": "notes",
    "E3": "scoop_size",
    "E4": "servings_per_unit",
    "E5": "quote_id",
    "E6": "jar_lid",
    "H9": "servings_base",
}


class RecordSheetOverflow(ValueError):
    """A section has more lines than the template's fixed rows hold (growing it would break its SUMs)."""


@dataclass
class RecordLine:
    """One ingredient row as printed; `is_active` picks Actives vs. Excipients, and part_num/price may be None."""

    name: str
    label_claim_mg: float | None = None
    activity: float | None = None
    overage: float | None = None
    part_num: str | None = None
    price_per_kilo: float | None = None
    is_active: bool = True


@dataclass
class RecordSheet:
    """Everything one Record Sheet prints: the header inputs plus the ingredient lines."""

    brand: str | None = None
    product: str | None = None
    flavor: str | None = None
    sample_code: str | None = None
    notes: str | None = None
    scoop_size: str | None = None
    servings_per_unit: float | None = None
    quote_id: str | None = None
    jar_lid: str | None = None
    servings_base: float | None = None
    lines: list[RecordLine] = field(default_factory=list)


def render(record: RecordSheet, template_path: Path | str = DEFAULT_TEMPLATE_PATH) -> bytes:
    """Return the filled workbook as .xlsx bytes; raises RecordSheetOverflow before writing anything."""
    actives = [l for l in record.lines if l.is_active]
    excipients = [l for l in record.lines if not l.is_active]
    if len(actives) > ACTIVE_CAPACITY:
        raise RecordSheetOverflow(
            f"{len(actives)} active lines exceed the template's {ACTIVE_CAPACITY} active rows"
        )
    if len(excipients) > EXCIPIENT_CAPACITY:
        raise RecordSheetOverflow(
            f"{len(excipients)} excipient lines exceed the template's "
            f"{EXCIPIENT_CAPACITY} excipient rows"
        )

    wb = load_workbook(template_path)
    ws = wb.worksheets[0]

    for cell, attr in _HEADER_CELLS.items():
        ws[cell] = getattr(record, attr)

    _fill_section(ws, ACTIVE_FIRST_ROW, ACTIVE_CAPACITY, actives)
    _fill_section(ws, EXCIPIENT_FIRST_ROW, EXCIPIENT_CAPACITY, excipients)

    wb.calculation.fullCalcOnLoad = True
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _fill_section(ws, first_row: int, capacity: int, lines: list[RecordLine]) -> None:
    """Clear a section's input cells (the template ships with an example), then write the lines."""
    for r in range(first_row, first_row + capacity):
        for col in _LINE_INPUT_COLUMNS:
            ws[f"{col}{r}"] = None

    for offset, line in enumerate(lines):
        r = first_row + offset
        ws[f"A{r}"] = line.part_num
        ws[f"B{r}"] = line.name
        ws[f"C{r}"] = line.label_claim_mg
        ws[f"D{r}"] = line.activity
        ws[f"E{r}"] = line.overage
        # A blank price stays blank — writing 0 hides a missing cost (finding 4).
        if line.price_per_kilo is not None:
            ws[f"I{r}"] = line.price_per_kilo
