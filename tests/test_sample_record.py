"""Tests for the Sample Record Sheet renderer (sample_record_xlsx.py) and the
catalog resolver behind it (queries.record_line_catalog).

The renderer is F2c: a second rendering of the same flavor-profile data the
Flavor Sheet uses (SPEC "reuse the flavor profiles"). Pinned here:

  1. Inputs land in the mapped header and line cells; the two line sections
     (Actives, Flavor System and Excipients) fill from their first row down.
  2. The renderer fills, never rebuilds: the computed columns (Actual Input,
     g/run, Formula %, Cost/Unit, kg/run) and the totals keep the template's
     own formulas, and the template's example inputs are cleared.
  3. The template's fixed capacity is enforced — an over-long section is
     refused, not silently truncated (layout §1 finding 5).
  4. A missing price is left blank, never written as 0 (finding 4).
  5. Part Number and Price/kg resolve from the row a line references, so a
     repriced material flows through without touching the flavor profile.

The real template is company property and never committed (SPEC §6), so the
renderer runs against a synthetic workbook of the same shape, built below from
the layout map in docs/phase_f_templates_layout.md §1.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from dtf_materials import flavor_sheets as fs
from dtf_materials import queries
from dtf_materials.db import init_db
from dtf_materials.sample_record_xlsx import (
    RecordLine,
    RecordSheet,
    RecordSheetOverflow,
    render,
)

ACTIVE_FIRST, ACTIVE_LAST = 12, 34      # 23 rows
EXCIP_FIRST, EXCIP_LAST = 37, 48        # 12 rows


@pytest.fixture
def conn(tmp_path):
    """A connection to a fresh, empty database."""
    c = init_db(tmp_path / "test.db")
    yield c
    c.close()


@pytest.fixture
def template(tmp_path):
    """A synthetic workbook shaped like the Record Sheet template (layout §1), example inputs included."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Header block — labels in A/D, example values in B/E/H that a fill clears.
    for cell, value in (
        ("B3", "Example Brand"), ("B4", "Example Product"), ("B5", "Example Flavor"),
        ("B6", "EX260101-01"), ("B7", "example note"),
        ("E3", "10cc"), ("E4", 30), ("E5", "EX-010126"), ("E6", "Jar A"),
        ("H9", 25),
    ):
        ws[cell] = value

    def _write_line_row(r: int) -> None:
        """Write one row's example inputs and computed-column formulas."""
        # Example inputs (cleared by a fill) ...
        ws[f"A{r}"] = "DTF-EX"
        ws[f"B{r}"] = f"Example material {r}"
        ws[f"C{r}"] = 100
        ws[f"D{r}"] = 1
        ws[f"E{r}"] = 0
        ws[f"I{r}"] = 9.99
        # ... and the computed columns the renderer must leave alone.
        ws[f"F{r}"] = f"=IFERROR(C{r}*(1+E{r})/D{r},0)"
        ws[f"G{r}"] = f"=(F{r}*$H$9)/1000"
        ws[f"H{r}"] = f"=IFERROR((C{r}/$F$50%)/100,0)"
        ws[f"J{r}"] = f"=(F{r}/1000000)*I{r}*$E$4"
        ws[f"K{r}"] = f"=G{r}/1000"

    for r in range(ACTIVE_FIRST, ACTIVE_LAST + 1):
        _write_line_row(r)
    for r in range(EXCIP_FIRST, EXCIP_LAST + 1):
        _write_line_row(r)

    ws["J35"] = "=SUM(J12:J34)"
    ws["J49"] = "=SUM(J37:J48)"
    ws["F50"] = "=SUM(F12:F48)"
    ws["J50"] = "=J35+J49"

    path = tmp_path / "record_template.xlsx"
    wb.save(path)
    return path


def _open(data: bytes):
    """Load rendered .xlsx bytes and return its sheet."""
    return load_workbook(BytesIO(data))["Sheet1"]


def _sheet(**overrides) -> RecordSheet:
    """Return a RecordSheet with a full header, plus overrides."""
    base = dict(
        brand="Acme", product="Pre-Workout", flavor="Berry", sample_code="AC260923-01",
        notes="fine powder", scoop_size="11cc", servings_per_unit=30,
        quote_id="AC-092326", jar_lid="Jar B", servings_base=100,
    )
    base.update(overrides)
    return RecordSheet(**base)


# --- header + line placement -------------------------------------------------

def test_header_and_lines_land_in_the_mapped_cells(template):
    record = _sheet(lines=[
        RecordLine("Caffeine", label_claim_mg=200, activity=1.0, overage=0.1,
                   part_num="DTF-100", price_per_kilo=42.5, is_active=True),
        RecordLine("Maltodextrin", label_claim_mg=8000, activity=1.0, overage=0.0,
                   part_num="DTF-900", price_per_kilo=3.25, is_active=False),
    ])
    ws = _open(render(record, template))

    assert ws["B3"].value == "Acme" and ws["B4"].value == "Pre-Workout"
    assert ws["B5"].value == "Berry" and ws["B6"].value == "AC260923-01"
    assert ws["E3"].value == "11cc" and ws["E4"].value == 30
    assert ws["E5"].value == "AC-092326" and ws["E6"].value == "Jar B"
    assert ws["H9"].value == 100

    # First active line -> row 12; first excipient line -> row 37.
    assert [ws[f"{c}12"].value for c in "ABCDEI"] == ["DTF-100", "Caffeine", 200, 1.0, 0.1, 42.5]
    assert [ws[f"{c}37"].value for c in "ABCDEI"] == ["DTF-900", "Maltodextrin", 8000, 1.0, 0.0, 3.25]


def test_the_computed_columns_and_totals_keep_the_templates_formulas(template):
    record = _sheet(lines=[RecordLine("Caffeine", 200, 1.0, 0.1, "DTF-100", 42.5, True)])
    ws = _open(render(record, template))

    # The renderer wrote row 12's inputs but must not touch its formulas.
    assert ws["F12"].value == "=IFERROR(C12*(1+E12)/D12,0)"
    assert ws["G12"].value == "=(F12*$H$9)/1000"
    assert ws["H12"].value == "=IFERROR((C12/$F$50%)/100,0)"
    assert ws["J12"].value == "=(F12/1000000)*I12*$E$4"
    assert ws["K12"].value == "=G12/1000"
    # Totals untouched.
    assert ws["J35"].value == "=SUM(J12:J34)" and ws["J50"].value == "=J35+J49"


def test_the_templates_example_inputs_are_cleared(template):
    record = _sheet(lines=[RecordLine("Caffeine", 200, 1.0, 0.1, "DTF-100", 42.5, True)])
    ws = _open(render(record, template))

    # Rows past the single active line, and the whole excipient section, keep
    # no example inputs — but still keep their formulas.
    assert [ws[f"{c}13"].value for c in "ABCDEI"] == [None] * 6
    assert ws["F13"].value == "=IFERROR(C13*(1+E13)/D13,0)"
    assert [ws[f"{c}37"].value for c in "ABCDEI"] == [None] * 6


def test_a_blank_price_stays_blank_not_zero(template):
    # A lab-only flavor with no price: I must be empty, not 0 (finding 4).
    record = _sheet(lines=[RecordLine("Peach E00000001", 300, 1.0, 0.0, None, None, False)])
    ws = _open(render(record, template))
    assert ws["A37"].value is None            # no Part # either
    assert ws["I37"].value is None


def test_an_absent_header_field_is_cleared(template):
    record = _sheet(scoop_size=None, servings_per_unit=None, lines=[])
    ws = _open(render(record, template))
    assert ws["E3"].value is None and ws["E4"].value is None


def test_the_workbook_recalculates_on_open(template):
    assert load_workbook(BytesIO(render(_sheet(lines=[]), template))).calculation.fullCalcOnLoad


# --- fixed capacity (finding 5) ----------------------------------------------

def test_too_many_active_lines_is_refused(template):
    record = _sheet(lines=[RecordLine(f"A{i}", 1, 1, 0, is_active=True) for i in range(24)])
    with pytest.raises(RecordSheetOverflow):
        render(record, template)


def test_too_many_excipient_lines_is_refused(template):
    record = _sheet(lines=[RecordLine(f"E{i}", 1, 1, 0, is_active=False) for i in range(13)])
    with pytest.raises(RecordSheetOverflow):
        render(record, template)


def test_a_full_section_still_renders(template):
    record = _sheet(lines=[RecordLine(f"A{i}", 1, 1, 0, is_active=True) for i in range(23)])
    ws = _open(render(record, template))
    assert ws[f"B{ACTIVE_LAST}"].value == "A22"   # the 23rd line fills the last row


# --- reuse: part # and price resolve from the referenced row -----------------

def _seed_material(conn, name="Caffeine", part="DTF-100", price=42.5):
    """Insert a priced warehouse material and return its id."""
    cur = conn.execute(
        "INSERT INTO materials (dtf_part_num, material_name, current_price_per_kilo) VALUES (?,?,?)",
        (part, name, price),
    )
    conn.commit()
    return cur.lastrowid


def test_catalog_fields_resolve_from_the_referenced_row(conn):
    material_id = _seed_material(conn)
    conn.execute(
        "INSERT INTO lab_samples (rd_id, flavor_name, sample_code, dtf_part_num, price_per_kilo) "
        "VALUES ('RD-0001','Peach','E00000001',NULL,7.5)"
    )
    conn.commit()

    assert queries.record_line_catalog(conn, material_id=material_id) == {
        "part_num": "DTF-100", "price_per_kilo": 42.5,
    }
    # Lab-only sample: no Part #, has a price.
    assert queries.record_line_catalog(conn, rd_id="RD-0001") == {
        "part_num": None, "price_per_kilo": 7.5,
    }
    # A typed line references nothing.
    assert queries.record_line_catalog(conn) == {"part_num": None, "price_per_kilo": None}


def test_a_reprice_reaches_a_rendered_sheet_without_touching_the_profile(conn, template):
    """The whole point of referencing rows: build a flavor profile, render its
    record sheet, reprice the material, render again — the new price shows up
    though nothing on the profile changed."""
    material_id = _seed_material(conn, price=42.5)
    sheet = fs.create_sheet(conn, customer="Acme", product="Pre-Workout", servings=4)
    profile = fs.add_profile(conn, sheet, "Berry", "AC260923-01", base_mg=8000)
    fs.add_line(conn, profile, material_id=material_id, mg_per_serving=200)

    def _render() -> object:
        """Render the profile with catalog fields resolved now, and return its sheet."""
        lines = []
        for l in fs.get_lines(conn, profile):
            cat = queries.record_line_catalog(conn, material_id=l["material_id"], rd_id=l["rd_id"])
            lines.append(RecordLine(
                name=l["name"], label_claim_mg=l["mg_per_serving"],
                activity=1.0, overage=0.1, part_num=cat["part_num"],
                price_per_kilo=cat["price_per_kilo"], is_active=True,
            ))
        return _open(render(_sheet(lines=lines), template))

    assert _render()["I12"].value == 42.5
    conn.execute("UPDATE materials SET current_price_per_kilo = 55.0 WHERE material_id = ?", (material_id,))
    conn.commit()
    assert _render()["I12"].value == 55.0
