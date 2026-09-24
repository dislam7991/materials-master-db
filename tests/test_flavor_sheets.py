"""Tests for the Flavor Sheet: the data object (flavor_sheets.py) and the xlsx
renderer (flavor_sheet_xlsx.py).

Pinned here:

  1. A line has exactly one source — warehouse material, lab sample, or typed
     name — and the printed name is resolved from that source at read time.
  2. The loaders never touch flavor-sheet tables (authored data no loader can
     rebuild, same rule as formulas).
  3. The renderer fills the template rather than rebuilding it: inputs land in
     the mapped cells, the g column is the template's own formula, the
     template's example data is cleared, and a long flavor or a fifth flavor
     grows the form instead of being cut off.

The real template is company property and never committed (SPEC §6), so the
renderer runs against a synthetic workbook of the same shape, built below from
the layout map in docs/phase_f_templates_layout.md §2.
"""

from __future__ import annotations

import sqlite3
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side

from dtf_materials import etl
from dtf_materials import flavor_sheets as fs
from dtf_materials.db import PROJECT_ROOT, init_db
from dtf_materials.flavor_sheet_xlsx import Flavor, render
from dtf_materials.sources import CsvInventorySource

SYNTHETIC_CSV = PROJECT_ROOT / "data" / "synthetic" / "raw_material_inventory.csv"


@pytest.fixture
def conn(tmp_path):
    """A connection to a fresh, empty database."""
    c = init_db(tmp_path / "test.db")
    yield c
    c.close()


def _seed_material(conn, name="Citric Acid", part_num="DTF-100"):
    """Insert a warehouse material and return its id."""
    cur = conn.execute(
        "INSERT INTO materials (dtf_part_num, material_name) VALUES (?,?)", (part_num, name)
    )
    conn.commit()
    return cur.lastrowid


def _seed_lab_sample(conn, rd_id="RD-0001", flavor="Peach", code="E00000001"):
    """Insert a lab sample and return its RD-ID."""
    conn.execute(
        "INSERT INTO lab_samples (rd_id, flavor_name, sample_code) VALUES (?,?,?)",
        (rd_id, flavor, code),
    )
    conn.commit()
    return rd_id


# --- the data object ---------------------------------------------------------

def test_lines_resolve_their_name_from_each_source(conn):
    material_id = _seed_material(conn, "DL-Malic Acid")
    rd_id = _seed_lab_sample(conn, flavor="Peach", code="E00000001")
    sheet = fs.create_sheet(conn, customer="Acme", product="Pre-Workout", quote_id="Q-1", servings=4)
    profile = fs.add_profile(conn, sheet, "Peach", "X260922-01", base_mg=8400)

    fs.add_line(conn, profile, material_id=material_id, mg_per_serving=1200)
    fs.add_line(conn, profile, rd_id=rd_id, mg_per_serving=300)
    fs.add_line(conn, profile, typed_name="New Masking Agent", mg_per_serving=100)

    lines = fs.get_lines(conn, profile)
    assert [(l["name"], l["source"], l["mg_per_serving"]) for l in lines] == [
        ("DL-Malic Acid", "Warehouse", 1200),
        ("Peach E00000001", "Lab", 300),
        ("New Masking Agent", "Typed", 100),
    ]


def test_a_renamed_material_reaches_existing_lines(conn):
    material_id = _seed_material(conn, "Citric Acid")
    sheet = fs.create_sheet(conn, servings=4)
    profile = fs.add_profile(conn, sheet, "Lime")
    fs.add_line(conn, profile, material_id=material_id, mg_per_serving=800)

    conn.execute("UPDATE materials SET material_name = 'Citric Acid Anhydrous' WHERE material_id = ?", (material_id,))
    conn.commit()

    assert fs.get_lines(conn, profile)[0]["name"] == "Citric Acid Anhydrous"


@pytest.mark.parametrize("sources", [
    {},                                                   # no source
    {"typed_name": "X", "rd_id": "RD-0001"},              # two sources
])
def test_a_line_needs_exactly_one_source(conn, sources):
    _seed_lab_sample(conn)
    profile = fs.add_profile(conn, fs.create_sheet(conn), "Lime")
    with pytest.raises(sqlite3.IntegrityError):
        fs.add_line(conn, profile, **sources)


def test_a_line_cannot_point_at_a_missing_catalog_row(conn):
    profile = fs.add_profile(conn, fs.create_sheet(conn), "Lime")
    with pytest.raises(sqlite3.IntegrityError):
        fs.add_line(conn, profile, material_id=999)
    with pytest.raises(sqlite3.IntegrityError):
        fs.add_line(conn, profile, rd_id="RD-9999")


def test_copying_a_profile_brings_base_and_lines(conn):
    sheet = fs.create_sheet(conn, servings=4)
    first = fs.add_profile(conn, sheet, "Berry", base_mg=8400)
    fs.add_line(conn, first, typed_name="Acid", mg_per_serving=1200)
    fs.add_line(conn, first, typed_name="Sweetener", mg_per_serving=150)

    second = fs.add_profile(conn, sheet, "Lime", copy_lines_from=first)

    assert fs.get_profiles(conn, sheet)[1]["base_mg"] == 8400
    assert [(l["name"], l["mg_per_serving"]) for l in fs.get_lines(conn, second)] == [
        ("Acid", 1200), ("Sweetener", 150),
    ]
    # A copy, not a shared reference: editing the new one leaves the first alone.
    fs.update_line(conn, fs.get_lines(conn, second)[0]["flavor_profile_line_id"],
                   mg_per_serving=1500, position=1)
    assert fs.get_lines(conn, first)[0]["mg_per_serving"] == 1200


def test_deleting_a_profile_closes_the_gap(conn):
    sheet = fs.create_sheet(conn)
    a, b, c = (fs.add_profile(conn, sheet, n) for n in ("A", "B", "C"))
    fs.delete_profile(conn, b)
    assert [(p["flavor_name"], p["position"]) for p in fs.get_profiles(conn, sheet)] == [("A", 1), ("C", 2)]


def test_deleting_a_sheet_removes_its_profiles_and_lines(conn):
    sheet = fs.create_sheet(conn)
    profile = fs.add_profile(conn, sheet, "A")
    fs.add_line(conn, profile, typed_name="Acid")
    fs.delete_sheet(conn, sheet)
    assert conn.execute("SELECT COUNT(*) FROM flavor_profiles").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM flavor_profile_lines").fetchone()[0] == 0


def test_the_etl_leaves_flavor_sheets_untouched(tmp_path):
    db_path = tmp_path / "etl.db"
    etl.run(CsvInventorySource(SYNTHETIC_CSV), db_path)
    conn = init_db(db_path)
    material_id = conn.execute("SELECT material_id FROM materials LIMIT 1").fetchone()[0]
    sheet = fs.create_sheet(conn, customer="Acme", servings=4)
    profile = fs.add_profile(conn, sheet, "Berry", base_mg=8400)
    fs.add_line(conn, profile, material_id=material_id, mg_per_serving=1200)
    before = fs.get_lines(conn, profile)
    conn.close()

    etl.run(CsvInventorySource(SYNTHETIC_CSV), db_path)

    conn = init_db(db_path)
    assert fs.get_sheet(conn, sheet)["customer"] == "Acme"
    assert [tuple(r) for r in fs.get_lines(conn, profile)] == [tuple(r) for r in before]
    conn.close()


def test_lines_resolve_price_from_each_source_and_a_reprice_flows_through(conn):
    material_id = _seed_material(conn, "Caffeine")
    conn.execute(
        "UPDATE materials SET current_price_per_kilo = 42.5 WHERE material_id = ?", (material_id,)
    )
    rd_id = _seed_lab_sample(conn)
    conn.execute("UPDATE lab_samples SET price_per_kilo = 7.5 WHERE rd_id = ?", (rd_id,))
    conn.commit()
    profile = fs.add_profile(conn, fs.create_sheet(conn), "Peach")
    fs.add_line(conn, profile, material_id=material_id, mg_per_serving=200)
    fs.add_line(conn, profile, rd_id=rd_id, mg_per_serving=300)
    fs.add_line(conn, profile, typed_name="Masking", mg_per_serving=100)

    assert [l["price_per_kilo"] for l in fs.get_lines(conn, profile)] == [42.5, 7.5, None]

    # Referenced, not copied: a reprice reaches the existing line.
    conn.execute(
        "UPDATE materials SET current_price_per_kilo = 55.0 WHERE material_id = ?", (material_id,)
    )
    conn.commit()
    assert fs.get_lines(conn, profile)[0]["price_per_kilo"] == 55.0


def test_profile_cost_sums_priced_lines_and_flags_the_rest(conn):
    material_id = _seed_material(conn, "Caffeine")
    conn.execute(
        "UPDATE materials SET current_price_per_kilo = 50 WHERE material_id = ?", (material_id,)
    )
    conn.commit()
    profile = fs.add_profile(conn, fs.create_sheet(conn, servings=4), "Berry")
    fs.add_line(conn, profile, material_id=material_id, mg_per_serving=200)  # 200/1e6*50 = 0.01/serving
    fs.add_line(conn, profile, typed_name="No price", mg_per_serving=100)   # priced? no -> unpriced
    fs.add_line(conn, profile, typed_name="No amount yet")                  # no mg -> not a missing price

    cost = fs.profile_cost(fs.get_lines(conn, profile), servings=4)
    assert cost["priced"] == 1 and cost["unpriced"] == 1
    assert cost["per_serving"] == pytest.approx(0.01)
    assert cost["per_sample"] == pytest.approx(0.04)


def test_profile_cost_is_none_when_nothing_prices(conn):
    profile = fs.add_profile(conn, fs.create_sheet(conn, servings=4), "Berry")
    fs.add_line(conn, profile, typed_name="No price", mg_per_serving=100)
    cost = fs.profile_cost(fs.get_lines(conn, profile), servings=4)
    assert cost["per_serving"] is None and cost["per_sample"] is None
    assert cost["priced"] == 0 and cost["unpriced"] == 1


def test_cost_helpers():
    assert fs.line_cost_per_serving(200, 50) == pytest.approx(0.01)
    assert fs.line_cost_per_serving(None, 50) is None
    assert fs.line_cost_per_serving(200, None) is None


def test_helpers():
    assert fs.product_line({"customer": "Acme", "product": " ", "quote_id": "Q-1"}) == "Acme - Q-1"
    assert fs.suggest_sample_id("SMPL", date(2026, 9, 16), 2) == "SMPL260916-02"
    assert fs.suggest_sample_id(None, date(2026, 9, 16), 1) == "260916-01"
    assert fs.grams_per_sample(8400, 4) == 33.6
    assert fs.grams_per_sample(None, 4) is None
    assert fs.grams_per_sample(100, None) is None


# --- the renderer ------------------------------------------------------------

BAND = PatternFill("solid", fgColor="D0CECE")
THIN, MEDIUM = Side(style="thin"), Side(style="medium")


@pytest.fixture
def template(tmp_path):
    """A synthetic workbook shaped like the real Flavor Sheet template, example data included.

    Two blocks of two slots (BASE + 12 rows top, BASE + 11 bottom), banding,
    thick bottom borders, and the g formula on the first eight rows of each slot.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Flavor Sheet"
    ws.merge_cells("A1:H2")
    ws["A3"] = "Product Name: Customer Name + Product Name + Quote ID"
    ws.merge_cells("A3:D4")
    ws["E3"], ws["F3"], ws["G3"], ws["H3"] = "Servings per flavor:", 6, "Serv. per retain:", 1
    ws["E4"], ws["F4"], ws["G4"] = "Servings per sample:", 4, "Scoop Size:"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.create_sheet("Sheet2")["A1"] = "older variant"

    for title_row, first, last in ((6, 8, 20), (21, 23, 34)):
        for title_col, name_col, mg_col, g_col in (("B", "A", "B", "C"), ("F", "E", "F", "G")):
            ws[f"{name_col}{title_row}"] = "Material"
            ws[f"{title_col}{title_row}"] = "Example Flavor + Sample ID"
            ws[f"{mg_col}{title_row + 1}"], ws[f"{g_col}{title_row + 1}"] = "mg", "g"
            ws[f"{name_col}{first}"] = "BASE"
            ws[f"{mg_col}{first}"] = 8400
            for r in range(first, first + 8):
                if r > first:
                    ws[f"{name_col}{r}"] = f"Example {r}"
                    ws[f"{mg_col}{r}"] = 100
                ws[f"{g_col}{r}"] = f"={mg_col}{r}*$F$3/1000"
        ws.merge_cells(f"B{title_row}:D{title_row}")
        ws.merge_cells(f"F{title_row}:H{title_row}")
        for r in range(first, last + 1):
            for c in range(1, 9):
                cell = ws.cell(r, c)
                cell.border = Border(bottom=MEDIUM if r == last else THIN)
                if (r - first) % 2:
                    cell.fill = BAND
        ws.row_dimensions[last].height = 15

    path = tmp_path / "template.xlsx"
    wb.save(path)
    return path


def _open(data: bytes):
    """Load rendered .xlsx bytes as a workbook."""
    return load_workbook(BytesIO(data))


def test_render_writes_header_and_lines_into_the_template(template):
    flavors = [
        Flavor("Berry X-01", 8400, [("Acid", 1200), ("Color", None)]),
        Flavor("Lime X-02", 8300, [("Acid", 1500)]),
    ]
    ws = _open(render("Acme - Pre - Q-1", 4, flavors, template))["Sheet1"]

    assert ws["A3"].value == "Product Name: Acme - Pre - Q-1"
    assert ws["F3"].value == 4
    # Inputs the flavor sheet doesn't ask for are cleared, not left as placeholders.
    assert ws["H3"].value is None and ws["F4"].value is None
    assert (ws["B6"].value, ws["F6"].value) == ("Berry X-01", "Lime X-02")
    assert [ws[f"{c}8"].value for c in "ABC"] == ["BASE", 8400, "=B8*$F$3/1000"]
    assert [ws[f"{c}9"].value for c in "ABC"] == ["Acid", 1200, "=B9*$F$3/1000"]
    assert [ws[f"{c}9"].value for c in "EFG"] == ["Acid", 1500, "=F9*$F$3/1000"]
    # Blank mg: no formula, so no invented 0.000 g.
    assert [ws[f"{c}10"].value for c in "ABC"] == ["Color", None, None]
    # The template's example data is gone, including the unused bottom slots.
    assert ws["A11"].value is None and ws["E10"].value is None
    assert ws["B21"].value is None and ws["A24"].value is None and ws["B23"].value is None
    assert ws["A23"].value == "BASE"   # static label stays


def test_a_long_flavor_grows_its_block_and_pushes_the_rest_down(template):
    long_lines = [(f"Material {i}", 10 * i) for i in range(1, 16)]   # BASE + 15 > 13 rows
    flavors = [Flavor("Long", 8400, long_lines), Flavor("Short", 8400, []), Flavor("Third", 8000, [("Acid", 1)])]
    ws = _open(render("P", 4, flavors, template))["Sheet1"]

    grown = 3   # 16 rows needed, 13 available
    assert ws["A23"].value == "Material 15" and ws["C23"].value == "=B23*$F$3/1000"
    assert ws["A23"].border.bottom.style == "medium"          # block still closes with its thick border
    assert ws.row_dimensions[20 + grown].height == 15
    # Banding carries on through the inserted rows.
    fills = [ws.cell(r, 1).fill.fill_type == "solid" for r in range(8, 24)]
    assert all(a != b for a, b in zip(fills, fills[1:]))
    # The bottom block moved down intact: title, merge and data.
    assert ws[f"B{21 + grown}"].value == "Third"
    assert f"B{21 + grown}:D{21 + grown}" in {str(m) for m in ws.merged_cells.ranges}
    assert [ws[f"{c}{24 + grown}"].value for c in "ABC"] == ["Acid", 1, f"=B{24 + grown}*$F$3/1000"]


def test_more_than_four_flavors_go_on_a_second_page(template):
    flavors = [Flavor(f"Flavor {i}", 8000, [("Acid", i)]) for i in range(1, 6)]
    wb = _open(render("P", 4, flavors, template))

    assert wb.sheetnames == ["Sheet1", "Flavors 5-8", "Sheet2"]
    page2 = wb["Flavors 5-8"]
    assert page2["B6"].value == "Flavor 5"
    assert page2["F6"].value is None and page2["B21"].value is None
    assert page2["A3"].value == "Product Name: P" and page2["F3"].value == 4
    assert wb["Sheet1"]["F21"].value == "Flavor 4"
    assert page2.sheet_properties.pageSetUpPr.fitToPage


def test_the_workbook_recalculates_on_open(template):
    assert _open(render("P", 4, [], template)).calculation.fullCalcOnLoad
